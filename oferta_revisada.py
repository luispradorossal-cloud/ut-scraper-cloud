"""
UT El Salvador - Oferta Revisada

Toma el Prog_Diaria_Inicial_Aislado que acaba de descargar scraper.py, le agrega
la columna "Precio Oferta" en la hoja Resumen y lo envia por correo como
"OFERTA REVISADA".

Regla: Precio Oferta = CMO * (1 - descuento). Descuento por defecto 8.5%.

El descuento queda en una celda editable del propio Excel (O1) y las 24 formulas
la referencian, asi que se puede cambiar el porcentaje en Excel sin tocar nada aca.

Variables de entorno:
    GMAIL_USER      - Correo Gmail remitente
    GMAIL_PASS      - Contrasena de aplicacion Gmail
    DEST_EMAIL      - Destinatarios (uno o varios separados por coma)
    DESCUENTO       - Opcional. Decimal, default 0.085
"""

import email.encoders
import email.mime.base
import email.mime.multipart
import email.mime.text
import os
import re
import smtplib
import sys
from copy import copy
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --- Config -------------------------------------------------------------------
TZ_SV = timezone(timedelta(hours=-6))

DOWNLOADS = Path("downloads")
PATRON_AISLADO = "Prog_Diaria_Inicial_Aislado_*.xlsx"

HOJA = "Resumen"
FILA_ENCABEZADO = 2
ENCABEZADO_NUEVO = "Precio Oferta [$/MWh]"
DESCUENTO_DEFECTO = 0.085

SUFIJO_SALIDA = "_OFERTA_REVISADA"


def now_sv():
    return datetime.now(TZ_SV)


def fecha_desde_nombre(nombre):
    """Prog_Diaria_Inicial_Aislado_060926.xlsx -> 06/09/2026"""
    m = re.search(r"(\d{6})\.xlsx$", nombre)
    if m:
        try:
            return datetime.strptime(m.group(1), "%d%m%y").strftime("%d/%m/%Y")
        except ValueError:
            pass
    return now_sv().strftime("%d/%m/%Y")


# --- Excel --------------------------------------------------------------------

def encontrar_columna_cmo(ws):
    """Ubica la columna del CMO por su encabezado, no por posicion fija."""
    for col in range(1, ws.max_column + 1):
        valor = ws.cell(FILA_ENCABEZADO, col).value
        if valor and "CMO" in str(valor).upper():
            return col
    return None


def primera_columna_libre(ws):
    """Primera columna sin ningun dato en toda la hoja."""
    for c in range(1, ws.max_column + 2):
        if all(ws.cell(r, c).value is None for r in range(1, ws.max_row + 1)):
            return c
    return ws.max_column + 1


def filas_horarias(ws):
    """Filas del primer bloque horario (columna B con enteros 0-23)."""
    filas = []
    for r in range(FILA_ENCABEZADO + 1, ws.max_row + 1):
        hora = ws.cell(r, 2).value
        if isinstance(hora, (int, float)) and float(hora).is_integer() and 0 <= hora <= 23:
            if filas and r != filas[-1] + 1:
                break   # empezo otra tabla (cotas de embalses, caudales)
            filas.append(r)
    return filas


def agregar_columna_oferta(entrada, salida, descuento):
    wb = openpyxl.load_workbook(entrada)

    if HOJA not in wb.sheetnames:
        print(f"  [ERROR] El archivo no tiene la hoja '{HOJA}'. Hojas: {wb.sheetnames}")
        return None

    ws = wb[HOJA]

    col_cmo = encontrar_columna_cmo(ws)
    if col_cmo is None:
        print(f"  [ERROR] No se encontro encabezado con 'CMO' en la fila {FILA_ENCABEZADO}.")
        print("          Puede que el UT haya cambiado el formato del archivo.")
        return None

    filas = filas_horarias(ws)
    if not filas:
        print("  [ERROR] No se encontraron filas horarias (columna B con 0-23).")
        return None

    col_nueva = primera_columna_libre(ws)
    L = get_column_letter(col_cmo)
    N = get_column_letter(col_nueva)
    P = get_column_letter(col_nueva + 1)

    # Celda editable con el descuento
    ws.cell(1, col_nueva, "Descuento oferta").font = Font(bold=True)
    celda = ws.cell(1, col_nueva + 1, descuento)
    celda.number_format = "0.0%"
    celda.font = Font(bold=True)
    celda.fill = PatternFill("solid", fgColor="FFF2CC")

    # Encabezado, copiando el formato del encabezado de CMO
    enc_ref = ws.cell(FILA_ENCABEZADO, col_cmo)
    enc = ws.cell(FILA_ENCABEZADO, col_nueva, ENCABEZADO_NUEVO)
    enc.font = Font(bold=True)
    enc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    try:
        if enc_ref.fill and enc_ref.fill.fgColor and enc_ref.fill.fgColor.rgb:
            enc.fill = PatternFill("solid", fgColor=enc_ref.fill.fgColor.rgb)
        enc.border = copy(enc_ref.border)
    except Exception:
        pass

    # Formulas hora por hora
    for r in filas:
        c = ws.cell(r, col_nueva, f"={L}{r}*(1-${P}$1)")
        c.number_format = "0.00"
        try:
            c.border = copy(ws.cell(r, col_cmo).border)
        except Exception:
            pass

    # Promedio del dia, si el archivo lo trae
    fila_prom = filas[-1] + 1
    if ws.cell(fila_prom, col_cmo).value is not None:
        c = ws.cell(fila_prom, col_nueva, f"={L}{fila_prom}*(1-${P}$1)")
        c.number_format = "0.00"
        c.font = Font(bold=True)

    ws.column_dimensions[N].width = 16
    ws.column_dimensions[P].width = 10

    wb.save(salida)

    print(f"  [OK] Columna CMO:      {L} ('{enc_ref.value}')")
    print(f"  [OK] Columna agregada: {N} ('{ENCABEZADO_NUEVO}')")
    print(f"  [OK] Descuento:        {descuento:.1%} (editable en {P}1)")
    print(f"  [OK] Formulas:         filas {filas[0]}-{filas[-1]} ({len(filas)} horas)")

    # Devolver el CMO promedio para ponerlo en el cuerpo del correo
    wb_val = openpyxl.load_workbook(entrada, data_only=True)
    cmo_prom = wb_val[HOJA].cell(fila_prom, col_cmo).value
    return cmo_prom


# --- Correo -------------------------------------------------------------------

def enviar(filepath, fecha_str, descuento, cmo_prom):
    gmail_user = os.environ["GMAIL_USER"]
    gmail_pass = os.environ["GMAIL_PASS"]
    dest_email = os.environ["DEST_EMAIL"]

    ahora = now_sv()

    msg = email.mime.multipart.MIMEMultipart()
    msg["From"] = gmail_user
    msg["To"] = dest_email
    msg["Subject"] = f"OFERTA REVISADA - {fecha_str}"

    if cmo_prom:
        linea_precio = (
            f"CMO promedio:      {cmo_prom:.2f} $/MWh\n"
            f"Oferta promedio:   {cmo_prom * (1 - descuento):.2f} $/MWh\n"
        )
    else:
        linea_precio = ""

    body = (
        f"OFERTA REVISADA - UT El Salvador\n"
        f"{'=' * 45}\n\n"
        f"Fecha de operacion: {fecha_str}\n"
        f"Generado:           {ahora.strftime('%d/%m/%Y %H:%M')} (hora SV)\n"
        f"Archivo:            {filepath.name} ({filepath.stat().st_size / 1024:.0f} KB)\n\n"
        f"Regla aplicada:     Precio Oferta = CMO x (1 - {descuento:.1%})\n"
        f"{linea_precio}\n"
        f"La columna nueva esta en la hoja 'Resumen'. El porcentaje de descuento\n"
        f"vive en una celda amarilla arriba de esa columna: cambialo ahi y las 24\n"
        f"horas se recalculan solas.\n\n"
        f"---\n"
        f"Enviado automaticamente por GitHub Actions\n"
    )
    msg.attach(email.mime.text.MIMEText(body, "plain", "utf-8"))

    with open(filepath, "rb") as f:
        adj = email.mime.base.MIMEBase("application", "octet-stream")
        adj.set_payload(f.read())
    email.encoders.encode_base64(adj)
    adj.add_header("Content-Disposition", f"attachment; filename={filepath.name}")
    msg.attach(adj)

    print(f"  [CORREO] Enviando a {dest_email}...", end=" ")
    with smtplib.SMTP("smtp.gmail.com", 587) as server:
        server.starttls()
        server.login(gmail_user, gmail_pass)
        server.send_message(msg)
    print("OK!")


# --- Main ---------------------------------------------------------------------

def main():
    descuento = float(os.environ.get("DESCUENTO", DESCUENTO_DEFECTO))

    print("=" * 60)
    print("  OFERTA REVISADA - UT EL SALVADOR")
    print("=" * 60)

    candidatos = sorted(DOWNLOADS.glob(PATRON_AISLADO))
    if not candidatos:
        print(f"  [ERROR] No hay ningun archivo {PATRON_AISLADO} en {DOWNLOADS}/")
        print("          Este paso corre despues de scraper.py con TIPO=aislado.")
        return 1

    entrada = candidatos[-1]
    fecha_str = fecha_desde_nombre(entrada.name)
    salida = entrada.with_name(entrada.stem + SUFIJO_SALIDA + ".xlsx")

    print(f"  Entrada:  {entrada.name}")
    print(f"  Fecha:    {fecha_str}")
    print(f"  Salida:   {salida.name}")
    print("-" * 60)

    cmo_prom = agregar_columna_oferta(entrada, salida, descuento)
    if cmo_prom is None and not salida.exists():
        return 1

    print("-" * 60)
    enviar(salida, fecha_str, descuento, cmo_prom)

    print("=" * 60)
    print(f"  COMPLETADO - {now_sv().strftime('%H:%M:%S')} (hora SV)")
    print("=" * 60)
    return 0


if __name__ == "__main__":
    sys.exit(main())
