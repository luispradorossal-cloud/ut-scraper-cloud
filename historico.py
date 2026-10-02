"""
Descarga histórica de Programación Diaria UT (aislado + programa final)
y extracción a CSV compacto (una fila por fecha-hora-tipo).

Lee el rango de historico/request.json y escribe historico/data/ut_hourly_<año>.csv
y historico/data/manifest.csv. No guarda los .xlsx (solo los datos extraídos).
"""
import csv, io, json, re, sys, time
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from scraper import create_session, list_files, BASE_URL, DOWNLOAD_PARAMS_BASE, PORTLET_ID

OUT = Path("historico/data"); OUT.mkdir(parents=True, exist_ok=True)
REQ = json.loads(Path("historico/request.json").read_text())
START, END = date.fromisoformat(REQ["start"]), date.fromisoformat(REQ["end"])

RX = {
    "A": re.compile(r"^Prog_Diaria_Inicial_Aislado_(\d{6})(?:_(\d+))?\.xlsx$", re.I),
    "F": re.compile(r"^Prog_Diaria(\d{6})(?:_(\d+))?\.xlsx$", re.I),
}
EMB = ["guaj", "cgra", "5nov", "15se", "3feb"]
VA_UNITS = ["15se-g1", "5nov-u1", "cgra-g1", "guaj-u1", "3feb-g1", "nepo-g1", "edp-cc"]
FIELDS = (["fecha", "tipo", "archivo", "hora", "cmo", "hidro", "grnc", "geo", "biomasa", "gnl", "termica",
           "intercambio", "total", "vol_marginal", "nepo"]
          + [f"caudal_{e}" for e in EMB] + [f"cota_{e}" for e in EMB] + [f"cv_{u}" for u in VA_UNITS])

z = lambda v: 0.0 if v is None or isinstance(v, str) else float(v)


def download_bytes(session, filename, folder):
    params = DOWNLOAD_PARAMS_BASE.copy()
    params["p_p_resource_id"] = filename
    params[f"_{PORTLET_ID}_folder"] = folder
    url = BASE_URL + "?" + "&".join(f"{k}={v}" for k, v in params.items())
    for attempt in range(3):
        try:
            r = session.get(url, timeout=90)
            r.raise_for_status()
            if "text/html" in r.headers.get("Content-Type", ""):
                return None
            return r.content
        except Exception as e:
            print(f"  retry {attempt+1} {filename}: {e}")
            time.sleep(5)
    return None


def section_rows(rows, title_kw):
    for i, r in enumerate(rows):
        if r and r[0] and title_kw.lower() in str(r[0]).lower():
            return rows[i + 2:i + 26]
    return None


def extract(content, d, tipo, fname):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    res = [tuple(r) for r in wb["Resumen"].iter_rows(values_only=True)]
    hourly = res[2:26]
    cau = section_rows(res, "Caudales")
    cot = section_rows(res, "Cotas")
    cv = {}
    if "Lista de Mérito" in wb.sheetnames:
        for r in wb["Lista de Mérito"].iter_rows(min_row=3, values_only=True):
            if len(r) > 2 and r[1] and isinstance(r[2], (int, float)):
                cv[str(r[1])] = float(r[2])
    inj_h, inj = None, []
    if "Inyeccion" in wb.sheetnames:
        ir = [tuple(r) for r in wb["Inyeccion"].iter_rows(values_only=True)]
        inj_h, inj = list(ir[1]), ir[2:26]
    out = []
    for h in range(24):
        r = hourly[h] if h < len(hourly) else None
        if not r or len(r) < 12 or r[11] is None:
            continue
        cmo = z(r[11])
        vol = nep = ""
        if inj_h and h < len(inj):
            vol = round(sum(z(inj[h][j]) for j, u in enumerate(inj_h)
                            if u in cv and abs(cv[u] - cmo) < 0.01), 2)
            nep = z(inj[h][inj_h.index("nepo-g1")]) if "nepo-g1" in inj_h else 0.0
        row = dict(fecha=d.isoformat(), tipo=tipo, archivo=fname, hora=h, cmo=cmo,
                   hidro=z(r[2]), grnc=z(r[3]), geo=z(r[4]), biomasa=z(r[5]), gnl=z(r[6]), termica=z(r[7]),
                   intercambio=z(r[8]), total=z(r[10]), vol_marginal=vol, nepo=nep)
        for k, e in enumerate(EMB):
            row[f"caudal_{e}"] = z(cau[h][2 + k]) if cau and h < len(cau) and len(cau[h]) > 2 + k else ""
            row[f"cota_{e}"] = z(cot[h][2 + k]) if cot and h < len(cot) and len(cot[h]) > 2 + k else ""
        for u in VA_UNITS:
            row[f"cv_{u}"] = cv.get(u, "")
        out.append(row)
    wb.close()
    return out


def main():
    s = create_session()
    years = range(START.year, END.year + 1)
    best = {}  # (date, tipo) -> (suffix, fname, folder)
    for y in years:
        for f in list_files(s, str(y)):
            for tipo, rx in RX.items():
                m = rx.match(f)
                if not m:
                    continue
                dd = m.group(1)
                try:
                    d = date(2000 + int(dd[4:6]), int(dd[2:4]), int(dd[0:2]))
                except ValueError:
                    continue
                if not (START <= d <= END):
                    continue
                suf = int(m.group(2) or 0)
                k = (d, tipo)
                if k not in best or suf > best[k][0]:
                    best[k] = (suf, f, str(y))
        print(f"año {y}: acumulado {len(best)} archivos objetivo")
    writers, handles = {}, {}
    man = open(OUT / "manifest.csv", "w", newline="")
    mw = csv.writer(man); mw.writerow(["fecha", "tipo", "archivo", "carpeta", "estado", "filas"])
    for i, ((d, tipo), (_, fname, folder)) in enumerate(sorted(best.items())):
        content = download_bytes(s, fname, folder)
        if not content:
            mw.writerow([d, tipo, fname, folder, "descarga_fallida", 0]); continue
        try:
            rows = extract(content, d, tipo, fname)
        except Exception as e:
            mw.writerow([d, tipo, fname, folder, f"error:{type(e).__name__}", 0]); continue
        if d.year not in writers:
            handles[d.year] = open(OUT / f"ut_hourly_{d.year}.csv", "w", newline="")
            writers[d.year] = csv.DictWriter(handles[d.year], fieldnames=FIELDS)
            writers[d.year].writeheader()
        for r in rows:
            writers[d.year].writerow(r)
        mw.writerow([d, tipo, fname, folder, "ok", len(rows)])
        if i % 50 == 0:
            print(f"  {i}/{len(best)} {fname}"); sys.stdout.flush()
        time.sleep(0.3)
    for h in handles.values(): h.close()
    man.close()
    print("listo")


if __name__ == "__main__":
    main()
